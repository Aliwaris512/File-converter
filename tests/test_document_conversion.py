from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from converters.office_converter import OfficeConverter
from services.conversion_service import ConversionService


@pytest.mark.skipif(OfficeConverter.libreoffice_path() is None, reason="LibreOffice is not installed")
def test_docx_to_pdf_uses_local_libreoffice(tmp_path: Path) -> None:
    source = tmp_path / "document with spaces.docx"
    document = Document()
    document.add_paragraph("Local document conversion")
    document.save(source)

    output = Path(ConversionService.convert_file(source, "PDF"))

    assert output.exists()
    assert output.read_bytes().startswith(b"%PDF")


def test_xlsx_to_csv_and_csv_to_xlsx(tmp_path: Path) -> None:
    workbook_path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active.append(["name", "count"])
    workbook.active.append(["apples", 3])
    workbook.save(workbook_path)

    csv_path = Path(ConversionService.convert_file(workbook_path, "CSV"))
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.reader(source))
    assert rows == [["name", "count"], ["apples", "3"]]

    xlsx_path = Path(ConversionService.convert_file(csv_path, "XLSX"))
    converted = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        assert list(converted.active.values) == [("name", "count"), ("apples", "3")]
    finally:
        converted.close()


def test_missing_libreoffice_is_reported_clearly(monkeypatch, tmp_path: Path) -> None:
    source = tmp_path / "document.docx"
    source.touch()
    monkeypatch.setattr("services.dependency_check.OfficeConverter.libreoffice_path", lambda: None)

    with pytest.raises(RuntimeError, match="LibreOffice is required"):
        ConversionService.validate_conversion(source, "PDF")


def test_existing_image_output_is_not_overwritten(tmp_path: Path) -> None:
    from PIL import Image

    source = tmp_path / "photo.png"
    Image.new("RGB", (5, 5), (20, 40, 60)).save(source)
    first_output = tmp_path / "photo_converted.jpeg"
    first_output.write_bytes(b"keep this existing file")

    created = Path(ConversionService.convert_file(source, "JPEG"))

    assert created == tmp_path / "photo_converted_1.jpeg"
    assert first_output.read_bytes() == b"keep this existing file"
