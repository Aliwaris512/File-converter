from __future__ import annotations

import csv
import shutil
import subprocess
import tempfile
from pathlib import Path

class OfficeConverter:
    """Office-format conversions backed by local LibreOffice and openpyxl."""

    @staticmethod
    def libreoffice_path() -> str | None:
        for command in ("soffice", "libreoffice"):
            executable = shutil.which(command)
            if executable:
                return executable
        for candidate in (
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        ):
            if candidate.is_file():
                return str(candidate)
        return None

    @staticmethod
    def convert_to_pdf(input_path: str | Path, output_path: str | Path) -> str:
        source = Path(input_path)
        destination = Path(output_path)
        soffice = OfficeConverter.libreoffice_path()
        if not soffice:
            raise RuntimeError(
                "LibreOffice is required for this Office-to-PDF conversion. Install LibreOffice locally and retry."
            )
        if destination.exists():
            raise FileExistsError(f"Output already exists: {destination}")

        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="localfileconverter-") as temp_dir:
            command = [
                soffice,
                "--headless",
                f"-env:UserInstallation={Path(temp_dir, 'profile').as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(source),
            ]
            result = subprocess.run(command, capture_output=True, text=True, timeout=180)
            generated = Path(temp_dir) / f"{source.stem}.pdf"
            if result.returncode != 0 or not generated.is_file():
                detail = result.stderr.strip() or result.stdout.strip()
                raise RuntimeError(detail or f"LibreOffice could not convert {source.name} to PDF.")
            shutil.copy2(generated, destination)
        return str(destination)

    @staticmethod
    def xlsx_to_csv(input_path: str | Path, output_path: str | Path) -> str:
        from openpyxl import load_workbook

        source = Path(input_path)
        destination = Path(output_path)
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            with destination.open("w", encoding="utf-8-sig", newline="") as output:
                writer = csv.writer(output)
                writer.writerows(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        return str(destination)

    @staticmethod
    def csv_to_xlsx(input_path: str | Path, output_path: str | Path) -> str:
        from openpyxl import Workbook

        source = Path(input_path)
        destination = Path(output_path)
        workbook = Workbook()
        worksheet = workbook.active
        with source.open("r", encoding="utf-8-sig", newline="") as input_file:
            for row in csv.reader(input_file):
                worksheet.append(row)
        workbook.save(destination)
        return str(destination)
